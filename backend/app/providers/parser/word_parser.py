from docx import Document as PyDocument

from app.models.document import Document
from app.models.parsed_document import ParsedDocument

from app.providers.parser.base import BaseParser


class WordParser(BaseParser):

    @property
    def name(self) -> str:
        return "PythonDocx"

    @property
    def version(self) -> str:
        return "1.0"

    @property
    def supported_content_types(self) -> list[str]:
        return [
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ]

    async def parse(
        self,
        document: Document,
    ) -> ParsedDocument:
        path = self._get_document_path(
            document
        )

        doc = PyDocument(str(path))

        paragraphs: list[str] = []

        # Extract paragraphs
        for paragraph in doc.paragraphs:
            text = paragraph.text.strip()

            if text:
                paragraphs.append(text)

        # Extract tables
        for table in doc.tables:
            for row in table.rows:
                cells = [
                    cell.text.strip()
                    for cell in row.cells
                    if cell.text.strip()
                ]

                if cells:
                    paragraphs.append(" | ".join(cells))

        content = "\n\n".join(paragraphs)

        parser_metadata = {
            "paragraph_count": len(doc.paragraphs),
            "section_count": len(doc.sections),
            "table_count": len(doc.tables),
            "language": "",
        }

        return ParsedDocument(
            document_id=document.id,
            content=content,
            parser_name=self.name,
            parser_version=self.version,
            parser_metadata=parser_metadata,
        )